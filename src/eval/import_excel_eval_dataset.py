"""Import the reviewed legal evaluation worksheet into the active JSON dataset."""

import argparse
import json
import math
import re
import sys
from pathlib import Path
from typing import Any

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - exercised by the CLI
    raise SystemExit("openpyxl is required; install dependencies from requirements.txt") from exc


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_SHEET = "Luật LĐ"
METADATA_COLUMNS = {
    "category": ("category",),
    "source_page": ("source page", "page"),
    "source_section": ("source section", "section"),
    "required_keywords": ("required keywords", "key concepts", "required keywords / key concepts"),
    "retrieval_difficulty": ("retrieval difficulty", "difficulty"),
    "expected_retrieval_type": ("expected retrieval type", "retrieval type"),
    "is_answerable": ("is answerable", "answerable"),
    "expected_behavior": ("expected behavior",),
    "notes": ("notes", "note"),
}


def _normalize_header(value: Any) -> str:
    text = str(value or "").strip().casefold()
    text = re.sub(r"[^\w]+", " ", text, flags=re.UNICODE)
    return " ".join(text.split())


def _clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value).strip() or None


def _answerable(value: Any, testcase_id: Any) -> Any:
    value = _clean(value)
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    normalized = str(value).strip().casefold()
    if normalized in {"yes", "true", "1"}:
        return True
    if normalized in {"no", "false", "0"}:
        return False
    raise ValueError(f"Test case {testcase_id}: unrecognized Is Answerable value {value!r}")


def _find_column(headers: dict[str, int], *aliases: str) -> int | None:
    normalized_aliases = {_normalize_header(alias) for alias in aliases}
    for header, index in headers.items():
        if header in normalized_aliases:
            return index
    return None


def convert_workbook(input_path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet {sheet_name!r} not found. Available: {workbook.sheetnames}")
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(rows)
        except StopIteration:
            raise ValueError(f"Worksheet {sheet_name!r} is empty")
        headers = {_normalize_header(value): i for i, value in enumerate(raw_headers) if value is not None}
        required_columns = {
            "id": _find_column(headers, "test case id", "testcase id", "id"),
            "query": _find_column(headers, "query", "question"),
            "expected_answer": _find_column(headers, "expected answer", "expected response", "ground truth"),
        }
        missing = [name for name, index in required_columns.items() if index is None]
        if missing:
            raise ValueError(f"Worksheet is missing required column(s): {', '.join(missing)}")
        evidence_column = _find_column(headers, "evidence", "reference evidence")
        metadata_columns = {key: _find_column(headers, *aliases) for key, aliases in METADATA_COLUMNS.items()}
        dataset = []
        missing_rows = []
        for excel_row, row in enumerate(rows, start=2):
            if not any(_clean(value) is not None for value in row):
                continue
            testcase_id = _clean(row[required_columns["id"]])
            query = _clean(row[required_columns["query"]])
            answer = _clean(row[required_columns["expected_answer"]])
            if not testcase_id or not query or not answer:
                missing_rows.append((excel_row, testcase_id, not bool(query), not bool(answer)))
                continue
            evidence = _clean(row[evidence_column]) if evidence_column is not None else None
            metadata = {}
            for key, column in metadata_columns.items():
                if column is None:
                    continue
                value = row[column] if column < len(row) else None
                value = _answerable(value, testcase_id) if key == "is_answerable" else _clean(value)
                metadata[key] = value
            dataset.append({
                "id": str(testcase_id),
                "query": str(query),
                "expected_answer": str(answer),
                "reference_contexts": [str(evidence)] if evidence is not None else [],
                "metadata": metadata,
            })
        if missing_rows:
            examples = "; ".join(f"row {row}: id={tid!r}, missing query={q}, answer={a}" for row, tid, q, a in missing_rows[:5])
            raise ValueError(f"{len(missing_rows)} data row(s) lack required values; {examples}")
        ids = [item["id"] for item in dataset]
        duplicates = sorted({value for value in ids if ids.count(value) > 1})
        if duplicates:
            raise ValueError(f"Duplicate testcase IDs: {', '.join(duplicates)}")
        return dataset
    finally:
        workbook.close()


def validate_dataset(dataset: list[dict[str, Any]]) -> None:
    if not isinstance(dataset, list):
        raise ValueError("Converted dataset root must be an array")
    ids = []
    for index, item in enumerate(dataset, start=1):
        if not all(isinstance(item.get(field), str) and item[field].strip() for field in ("id", "query", "expected_answer")):
            raise ValueError(f"Converted testcase {index} has an empty required field")
        if not isinstance(item.get("reference_contexts"), list):
            raise ValueError(f"Testcase {item['id']} reference_contexts must be an array")
        ids.append(item["id"])
        json.dumps(item, allow_nan=False)
    if len(set(ids)) != len(ids):
        raise ValueError("Converted dataset contains duplicate IDs")


def main() -> None:
    parser = argparse.ArgumentParser(description="Import an Excel worksheet into the RAG evaluation dataset")
    parser.add_argument("--input", type=Path, required=True, help="Source .xlsx workbook")
    parser.add_argument("--sheet", default=DEFAULT_SHEET)
    parser.add_argument("--output", type=Path, default=ROOT / "data" / "eval_dataset.json")
    args = parser.parse_args()
    if not args.input.is_file():
        parser.error(f"Input workbook not found: {args.input}")
    try:
        dataset = convert_workbook(args.input, args.sheet)
        validate_dataset(dataset)
    except (OSError, ValueError) as exc:
        parser.exit(1, f"Import failed: {exc}\n")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="\n") as handle:
        json.dump(dataset, handle, ensure_ascii=False, indent=2, allow_nan=False)
        handle.write("\n")
    evidence_count = sum(bool(item["reference_contexts"]) for item in dataset)
    print(f"Worksheet        : {args.sheet}")
    print(f"Imported cases   : {len(dataset)}")
    print(f"Unique IDs       : {len({item['id'] for item in dataset})}")
    print("Missing query    : 0")
    print("Missing answer   : 0")
    print(f"With evidence    : {evidence_count}")
    print(f"Without evidence : {len(dataset) - evidence_count}")
    print(f"Output           : {args.output}")


if __name__ == "__main__":
    main()
